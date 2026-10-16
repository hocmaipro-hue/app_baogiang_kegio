# xuat_excel_bieu_1c.py
import os
from pkg_data_service import lay_du_lieu_pkg
from pkg_excel_helper import (
    clone_row_style,
    find_row_by_marker,
    find_row_in_cols,
    safe_set,
    unmerge_and_merge_exact,
)
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, Side

thin_side = Side(border_style='thin', color='000000')
thin_border = Border(
    top=thin_side, left=thin_side, right=thin_side, bottom=thin_side
)

FONT_SIZE = 13
font_b = Font(name='Times New Roman', size=FONT_SIZE, bold=True)
font_n = Font(name='Times New Roman', size=FONT_SIZE, bold=False)
font_it = Font(name='Times New Roman', size=FONT_SIZE, italic=True)
al_c = Alignment(horizontal='center', vertical='center')
al_l = Alignment(horizontal='left', vertical='center')


def apply_border_range(ws, min_r, max_r, min_c, max_c):
    for r in range(min_r, max_r + 1):
        for c in range(min_c, max_c + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border


def ghi_du_lieu_2_cot(ws, dic_data, sec_h, next_h, is_nghirieng=False):
    f_row = find_row_in_cols(ws, sec_h, cols=('A', 'B'), min_r=15, max_r=250)
    if not f_row or not dic_data:
        if f_row:
            safe_set(ws, f_row, 9, 0, font=font_b, align=al_c)
        return 0.0

    f_next = find_row_in_cols(
        ws, next_h, cols=('A', 'B'), min_r=f_row + 1, max_r=250
    )
    if f_next and f_next > f_row + 1:
        ws.delete_rows(f_row + 1, f_next - f_row - 1)

    f_next = find_row_in_cols(
        ws, next_h, cols=('A', 'B'), min_r=f_row + 1, max_r=250
    )
    ins_row = f_next if f_next else f_row + 1
    dates = list(dic_data.keys())
    sub_total = 0.0

    for n in range(0, len(dates), 2):
        dL = dic_data[dates[n]]
        dR = dic_data[dates[n + 1]] if (n + 1 < len(dates)) else None
        rMax = max(len(dL['items']), len(dR['items']) if dR else 0)
        bR = (2 if is_nghirieng else 1) + rMax
        ws.insert_rows(ins_row, amount=bR)
        for off in range(bR):
            clone_row_style(ws, f_row, ins_row + off)

        r = ins_row

        if is_nghirieng:
            ws.row_dimensions[r].height = 22
            ly_do_L = (dL.get('ly_do', '') or '').strip()
            text_lydo_L = (
                f'- Lý do nghỉ:  {ly_do_L}' if ly_do_L else '- Lý do nghỉ:'
            )
            unmerge_and_merge_exact(
                ws, f'B{r}:G{r}', text_lydo_L, font=font_n, align=al_l
            )

            if dR:
                ly_do_R = (dR.get('ly_do', '') or '').strip()
                text_lydo_R = (
                    f'- Lý do nghỉ:  {ly_do_R}' if ly_do_R else '- Lý do nghỉ:'
                )
                unmerge_and_merge_exact(
                    ws, f'H{r}:I{r}', text_lydo_R, font=font_n, align=al_l
                )
            else:
                unmerge_and_merge_exact(
                    ws, f'H{r}:I{r}', '', font=font_n, align=al_l
                )
            r += 1

        d_str_L = (
            dates[n].strftime('%d/%m/%Y')
            if hasattr(dates[n], 'strftime')
            else str(dates[n])
        )
        ws.row_dimensions[r].height = 22
        safe_set(ws, r, 2, '- Ngày:', font=font_n, align=al_l)
        safe_set(ws, r, 3, d_str_L, font=font_n, align=al_c)
        if dR:
            d_str_R = (
                dates[n + 1].strftime('%d/%m/%Y')
                if hasattr(dates[n + 1], 'strftime')
                else str(dates[n + 1])
            )
            safe_set(ws, r, 7, '- Ngày:', font=font_n, align=al_l)
            safe_set(ws, r, 8, d_str_R, font=font_n, align=al_c)
        r += 1

        for idx, (lop, gio) in enumerate(dL['items']):
            cr = r + idx
            ws.row_dimensions[cr].height = 22
            safe_set(ws, cr, 2, '+ Lớp:', font=font_n, align=al_l)
            safe_set(ws, cr, 3, lop, font=font_n, align=al_c)
            safe_set(ws, cr, 4, gio, font=font_n, align=al_c)
            safe_set(ws, cr, 5, 'giờ', font=font_n, align=al_l)
            sub_total += gio

        if dR:
            for idx, (lop, gio) in enumerate(dR['items']):
                cr = r + idx
                safe_set(ws, cr, 7, '+ Lớp:', font=font_n, align=al_l)
                safe_set(ws, cr, 8, lop, font=font_n, align=al_c)
                safe_set(ws, cr, 9, gio, font=font_n, align=al_c)
                safe_set(ws, cr, 10, 'giờ', font=font_n, align=al_l)
                sub_total += gio
        ins_row += bR

    safe_set(ws, f_row, 9, sub_total, font=font_b, align=al_c)
    return sub_total


def tao_phieu_ke_khai_1c(
    session, tu_tuan, den_tuan, thang_str, extra_info=None
):
    magv = (
        session.get('magv', '')
        if isinstance(session, dict)
        else getattr(session, 'magv', '')
    )
    tengv = (
        session.get('tengv', 'Giáo viên')
        if isinstance(session, dict)
        else getattr(session, 'tengv', 'Giáo viên')
    )
    to_bm = (
        session.get('to_bo_mon', 'Toán')
        if isinstance(session, dict)
        else getattr(session, 'to_bo_mon', 'Toán')
    )

    D = lay_du_lieu_pkg(magv, tu_tuan, den_tuan, extra_info)

    template_name = (
        'Phong_NVH_PKG_t9.xlsx'
        if os.path.exists('Phong_NVH_PKG_t9.xlsx')
        else 'Mau_Bieu_1C.xlsx'
    )
    wb = openpyxl.load_workbook(template_name)
    ws = wb['PKG'] if 'PKG' in wb.sheetnames else wb.active

    # A. TIÊU ĐỀ
    safe_set(
        ws, 1, 1, 'Trường THPT Nguyễn Việt Hồng', font=font_b, align=al_l
    )
    safe_set(ws, 2, 1, f'Tổ: {to_bm}', font=font_b, align=al_l)

    unmerge_and_merge_exact(
        ws,
        'A5:K5',
        f"Tháng  {D['ngay_giua'].month:02d}  năm  {D['ngay_giua'].year}   Số"
        f" tuần:  {D['so_tuan']}",
        font=font_b,
        align=al_c,
    )
    unmerge_and_merge_exact(
        ws,
        'A6:K6',
        f"( Từ ngày  {D['d_start'].strftime('%d/%m/%Y')}  đến ngày "
        f" {D['d_end'].strftime('%d/%m/%Y')} )",
        font=font_it,
        align=al_c,
    )

    for r in range(7, 12):
        ws.row_dimensions[r].height = 22
    safe_set(ws, 8, 1, f'- Họ và tên: {tengv}', font=font_n, align=al_l)
    safe_set(ws, 9, 1, f"- Chức vụ: {D['chuc_vu']}", font=font_n, align=al_l)
    str_kn = f"- Các công tác kiêm nhiệm: {D['kn1']}" + (
        f", {D['kn2']}" if D['kn2'] else ''
    )
    safe_set(ws, 10, 1, str_kn, font=font_n, align=al_l)
    safe_set(
        ws,
        11,
        1,
        f"- Được phân công dạy môn, lớp: {D['str_phan_cong']}",
        font=font_n,
        align=al_l,
    )

    safe_set(ws, 12, 3, f"{D['dm_tuan']} tiết", font=font_n, align=al_c)
    safe_set(ws, 12, 5, f"{D['so_tuan']} tuần", font=font_n, align=al_c)
    safe_set(ws, 12, 9, D['tong_dm'], font=font_b, align=al_c)

    # B. MỤC 3: TKB
    r_start_tkb = 17
    r_sec4 = (
        find_row_in_cols(ws, '4. quá trình', cols=('A', 'B'), min_r=15, max_r=50)
        or 18
    )

    matrix = D['tkb_matrix']
    so_lop = len(matrix)
    tong_so_dong_can = max(so_lop, 1) + 1
    so_dong_co_san = r_sec4 - r_start_tkb

    if tong_so_dong_can > so_dong_co_san:
        ws.insert_rows(r_sec4, amount=tong_so_dong_can - so_dong_co_san)

    new_r_tong = r_start_tkb + so_lop

    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= r_start_tkb and rng.max_row <= new_r_tong + 1:
            ws.merged_cells.remove(rng)

    r_idx = r_start_tkb
    sum_tuan = [0] * 5

    for lop, info in matrix.items():
        ws.row_dimensions[r_idx].height = 22
        ws.cell(row=r_idx, column=2, value=lop).alignment = al_c
        ws.cell(row=r_idx, column=2).font = font_n

        t_lop = 0
        for w_idx in range(5):
            c_idx = w_idx + 3
            if w_idx < D['so_tuan']:
                val_tiet = info['weeks'][w_idx]
                ws.cell(
                    row=r_idx,
                    column=c_idx,
                    value=val_tiet if val_tiet > 0 else '',
                ).alignment = al_c
                ws.cell(row=r_idx, column=c_idx).font = font_n
                sum_tuan[w_idx] += val_tiet
                t_lop += val_tiet
            else:
                ws.cell(row=r_idx, column=c_idx, value='').alignment = al_c

        ws.cell(
            row=r_idx, column=8, value=t_lop if t_lop > 0 else ''
        ).alignment = al_c
        ws.cell(row=r_idx, column=8).font = font_b

        ws.cell(row=r_idx, column=9, value=info['notes']).alignment = al_l
        ws.cell(row=r_idx, column=9).font = font_n
        r_idx += 1

    # Dòng Tổng Mục 3
    ws.row_dimensions[new_r_tong].height = 22
    ws.cell(row=new_r_tong, column=2, value='Tổng').alignment = al_c
    ws.cell(row=new_r_tong, column=2).font = font_b

    for w_idx in range(5):
        c_idx = w_idx + 3
        val_w = sum_tuan[w_idx] if w_idx < D['so_tuan'] else ''
        ws.cell(row=new_r_tong, column=c_idx, value=val_w).alignment = al_c
        ws.cell(row=new_r_tong, column=c_idx).font = font_b

    sum2_tong = sum(sum_tuan)
    ws.cell(row=new_r_tong, column=8, value=sum2_tong).alignment = al_c
    ws.cell(row=new_r_tong, column=8).font = font_b
    ws.cell(row=new_r_tong, column=9, value='').alignment = al_c

    apply_border_range(ws, r_start_tkb, new_r_tong, 2, 9)
    safe_set(ws, 13, 9, sum2_tong, font=font_b, align=al_c)

    # C. MỤC 4.1: DẠY THAY
    r_41 = find_row_in_cols(
        ws, '4.1', cols=('A', 'B'), min_r=new_r_tong, max_r=70
    )
    sum3_dt = 0
    if r_41:
        r_hdr_41 = r_41 + 1
        ws.row_dimensions[r_hdr_41].height = 24
        unmerge_and_merge_exact(
            ws,
            f'A{r_hdr_41}:B{r_hdr_41}',
            'Ngày tháng năm',
            font=font_b,
            align=al_c,
        )
        unmerge_and_merge_exact(
            ws,
            f'C{r_hdr_41}:D{r_hdr_41}',
            'GV được dạy thay',
            font=font_b,
            align=al_c,
        )
        unmerge_and_merge_exact(
            ws,
            f'E{r_hdr_41}:H{r_hdr_41}',
            'Lí do phải dạy thay',
            font=font_b,
            align=al_c,
        )
        safe_set(ws, r_hdr_41, 9, 'Lớp', font=font_b, align=al_c)
        safe_set(ws, r_hdr_41, 10, 'Số tiết', font=font_b, align=al_c)
        apply_border_range(ws, r_hdr_41, r_hdr_41, 1, 10)

        r_data_41 = r_hdr_41 + 1
        so_dt = len(D['ds_day_thay'])

        if so_dt == 0:
            ws.row_dimensions[r_data_41].height = 22
            unmerge_and_merge_exact(
                ws, f'A{r_data_41}:B{r_data_41}', '', font=font_n, align=al_c
            )
            unmerge_and_merge_exact(
                ws, f'C{r_data_41}:D{r_data_41}', '', font=font_n, align=al_c
            )
            unmerge_and_merge_exact(
                ws, f'E{r_data_41}:H{r_data_41}', '', font=font_n, align=al_l
            )
            safe_set(ws, r_data_41, 9, '', font_n, al_c)
            safe_set(ws, r_data_41, 10, '', font_n, al_c)
            apply_border_range(ws, r_data_41, r_data_41, 1, 10)
        else:
            r_42 = find_row_in_cols(
                ws, '4.2', cols=('A', 'B'), min_r=r_41 + 1, max_r=80
            ) or (r_data_41 + 1)
            co_san = max(r_42 - r_data_41, 1)
            if so_dt > co_san:
                ws.insert_rows(r_42, amount=so_dt - co_san)

            for idx, dt in enumerate(D['ds_day_thay']):
                cr = r_data_41 + idx
                ws.row_dimensions[cr].height = 22
                unmerge_and_merge_exact(
                    ws,
                    f'A{cr}:B{cr}',
                    dt['ngay'].strftime('%d/%m/%Y'),
                    font=font_n,
                    align=al_c,
                )
                unmerge_and_merge_exact(
                    ws, f'C{cr}:D{cr}', dt['gv_thay'], font=font_n, align=al_c
                )
                unmerge_and_merge_exact(
                    ws, f'E{cr}:H{cr}', dt['ly_do'], font=font_n, align=al_l
                )
                safe_set(ws, cr, 9, dt['lop'], font_n, al_c)
                safe_set(ws, cr, 10, dt['tiet'], font_n, al_c)
                apply_border_range(ws, cr, cr, 1, 10)
                sum3_dt += dt['tiet']

        safe_set(ws, r_41, 9, sum3_dt, font=font_b, align=al_c)

    # MỤC 4.2
    r_42 = find_row_in_cols(ws, '4.2', cols=('A', 'B'), min_r=20, max_r=70)
    if r_42:
        safe_set(ws, r_42, 9, 0, font=font_b, align=al_c)

    # MỤC 4.3 (NGHỈ VIỆC RIÊNG)
    sum5_nghi = ghi_du_lieu_2_cot(
        ws, D['dic_nghi_rieng'], '4.3', '4.4', is_nghirieng=True
    )

    # MỤC 4.4 (ĐIỀU ĐỘNG)
    r_44 = find_row_in_cols(ws, '4.4', cols=('A', 'B'), min_r=30, max_r=100)
    sum_dd = 0
    if r_44:
        if D['ds_nghi_dieudong']:
            sum_dd = sum(x['tiet'] for x in D['ds_nghi_dieudong'])
            safe_set(ws, r_44, 9, sum_dd, font=font_b, align=al_c)
            dd0 = D['ds_nghi_dieudong'][0]
            safe_set(ws, r_44 + 1, 2, dd0['ly_do'], font=font_n, align=al_l)
            d_str = (
                dd0['ngay'].strftime('%d/%m/%Y')
                if hasattr(dd0['ngay'], 'strftime')
                else str(dd0['ngay'])
            )
            safe_set(ws, r_44 + 2, 2, f'- Ngày: {d_str}', font=font_n, align=al_l)
            safe_set(
                ws,
                r_44 + 3,
                2,
                f"+ Lớp: {dd0['lop']}       {sum_dd} giờ",
                font=font_n,
                align=al_l,
            )
        else:
            safe_set(ws, r_44, 9, 0, font=font_b, align=al_c)

    # MỤC 4.5 (DẠY BÙ)
    sum6_bu = ghi_du_lieu_2_cot(ws, D['dic_day_bu'], '4.5', '5.')

    # =========================================================================
    # E. MỤC 5: BDHSG (GIỮ CHUẨN 7.5, KHÔNG BỊ ÉP LÊN 8)
    # =========================================================================
    r_m7 = find_row_in_cols(ws, '5. số giờ', cols=('A', 'B'), min_r=35, max_r=120)
    X = float(D['tong_hsg'])

    if r_m7:
        r_sec6_cu = find_row_in_cols(
            ws, '6. số giờ', cols=('A', 'B'), min_r=r_m7 + 1, max_r=140
        )
        rB = r_m7 + 2

        gio_1, tuan_1, tong_1 = 0, 0, 0
        gio_2, tuan_2, tong_2 = 0, 0, 0

        if X > 0:
            if X % 1 != 0:
                con_lai = X - 1.5
                gio_2, tuan_2, tong_2 = 1.5, 1, 1.5
                if con_lai > 0:
                    gio_1, tuan_1, tong_1 = 3, int(con_lai // 3), con_lai
            elif X % 3 == 0:
                gio_1, tuan_1, tong_1 = 3, int(X // 3), int(X)
            else:
                for a in range(int(X // 3), -1, -1):
                    if (X - 3 * a) % 2 == 0:
                        b = int((X - 3 * a) // 2)
                        gio_1, tuan_1, tong_1 = 3, a, a * 3
                        gio_2, tuan_2, tong_2 = 2, b, b * 2
                        break

        so_dong_can = (1 if tong_1 > 0 else 0) + (1 if tong_2 > 0 else 0)

        if r_sec6_cu:
            so_dong_hien_co = r_sec6_cu - rB
            if so_dong_can > so_dong_hien_co:
                so_dong_chen = so_dong_can - so_dong_hien_co
                ws.insert_rows(r_sec6_cu, amount=so_dong_chen)
                for off in range(so_dong_chen):
                    clone_row_style(ws, rB, r_sec6_cu + off)

        for c in range(3, 9):
            safe_set(ws, rB, c, '')
            safe_set(ws, rB + 1, c, '')

        curr_fill_r = rB
        if tong_1 > 0:
            ws.row_dimensions[curr_fill_r].height = 22
            safe_set(ws, curr_fill_r, 3, gio_1, font=font_n, align=al_c)
            safe_set(ws, curr_fill_r, 4, 'giờ x', font=font_n, align=al_c)
            safe_set(ws, curr_fill_r, 5, tuan_1, font=font_n, align=al_c)
            safe_set(ws, curr_fill_r, 6, 'tuần =', font=font_n, align=al_c)
            safe_set(
                ws,
                curr_fill_r,
                7,
                int(tong_1)
                if isinstance(tong_1, float) and tong_1.is_integer()
                else tong_1,
                font=font_n,
                align=al_c,
            )
            safe_set(ws, curr_fill_r, 8, 'giờ', font=font_n, align=al_l)
            curr_fill_r += 1

        if tong_2 > 0:
            ws.row_dimensions[curr_fill_r].height = 22
            safe_set(ws, curr_fill_r, 3, gio_2, font=font_n, align=al_c)
            safe_set(ws, curr_fill_r, 4, 'giờ x', font=font_n, align=al_c)
            safe_set(ws, curr_fill_r, 5, tuan_2, font=font_n, align=al_c)
            safe_set(ws, curr_fill_r, 6, 'tuần =', font=font_n, align=al_c)
            safe_set(
                ws,
                curr_fill_r,
                7,
                int(tong_2)
                if isinstance(tong_2, float) and tong_2.is_integer()
                else tong_2,
                font=font_n,
                align=al_c,
            )
            safe_set(ws, curr_fill_r, 8, 'giờ', font=font_n, align=al_l)

        # 👉 GHI CHÍNH XÁC SỐ THỰC 7.5 VÀ ÉP FORMAT 0.0 ĐỂ KHÔNG BỊ TRÒN SỐ
        c_m7 = safe_set(ws, r_m7, 9, X, font=font_b, align=al_c)
        if X % 1 != 0:
            c_m7.number_format = '0.0'

    # =========================================================================
    # F. MỤC 6: GIẢM TRỪ KIÊM NHIỆM
    # =========================================================================
    r_m8 = find_row_in_cols(
        ws, '6. số giờ được giảm', cols=('A', 'B'), min_r=r_m7 + 2, max_r=160
    )

    sum8_kn = 0.0
    if r_m8:
        bR = r_m8
        ws.row_dimensions[bR].height = 24
        ws.row_dimensions[bR + 1].height = 24
        safe_set(
            ws,
            bR,
            1,
            '6. Số giờ được giảm giờ định mức đối với giáo viên kiêm nhiệm,',
            font=font_b,
            align=al_l,
        )
        safe_set(
            ws,
            bR + 1,
            1,
            'giảm giờ đối với các giáo viên hợp đồng lần đầu và giáo viên nữ có'
            ' con nhỏ dưới 12 tháng.',
            font=font_b,
            align=al_l,
        )

        if D['kn1'] or D['gt_kn1'] > 0:
            t = D['gt_kn1'] * D['so_tuan']
            safe_set(
                ws,
                bR + 2,
                2,
                f"+ Kiêm nhiệm công tác: {D['kn1']}",
                font=font_n,
                align=al_l,
            )
            safe_set(ws, bR + 3, 3, D['gt_kn1'], font=font_n, align=al_c)
            safe_set(ws, bR + 3, 5, D['so_tuan'], font=font_n, align=al_c)
            safe_set(ws, bR + 3, 7, t, font=font_n, align=al_c)
            sum8_kn += t

        if D['kn2'] or D['gt_kn2'] > 0:
            t = D['gt_kn2'] * D['so_tuan']
            safe_set(
                ws,
                bR + 4,
                2,
                f"+ Kiêm nhiệm công tác: {D['kn2']}",
                font=font_n,
                align=al_l,
            )
            safe_set(ws, bR + 5, 3, D['gt_kn2'], font=font_n, align=al_c)
            safe_set(ws, bR + 5, 5, D['so_tuan'], font=font_n, align=al_c)
            safe_set(ws, bR + 5, 7, t, font=font_n, align=al_c)
            sum8_kn += t

        if D['gt_khac'] > 0:
            t = D['gt_khac'] * D['so_tuan']
            safe_set(ws, bR + 9, 3, D['gt_khac'], font=font_n, align=al_c)
            safe_set(ws, bR + 9, 5, D['so_tuan'], font=font_n, align=al_c)
            safe_set(ws, bR + 9, 7, t, font=font_n, align=al_c)
            sum8_kn += t

        safe_set(ws, bR, 9, sum8_kn, font=font_b, align=al_c)

    # =========================================================================
    # G. MỤC 7 & 8: TỔNG KẾT (TÍNH CHUẨN SỐ LẺ .5)
    # =========================================================================
    r_m9 = find_row_in_cols(
        ws, '7. tổng số giờ', cols=('A', 'B'), min_r=r_m8 + 5, max_r=200
    )
    if not r_m9:
        r_m9 = r_m8 + 9
    r_sec8 = r_m9 + 1

    ws.row_dimensions[r_m9].height = 24
    ws.row_dimensions[r_sec8].height = 24

    safe_set(
        ws,
        r_m9,
        1,
        '7. Tổng số giờ thực hiện = (2)+(3)-(5)-(5a)+(6)+(7)+(8)',
        font=font_b,
        align=al_l,
    )
    safe_set(
        ws, r_sec8, 1, '8. Số giờ dạy thêm  = (9)-(1)', font=font_b, align=al_l
    )

    sum9 = (
        sum2_tong
        + sum3_dt
        - sum5_nghi
        - sum_dd
        + sum6_bu
        + float(X)
        + sum8_kn
    )
    c_m9 = safe_set(ws, r_m9, 9, sum9, font=font_b, align=al_c)
    if isinstance(sum9, float) and sum9 % 1 != 0:
        c_m9.number_format = '0.0'

    thua_gio = sum9 - D['tong_dm']
    c_sec8 = safe_set(ws, r_sec8, 9, thua_gio, font=font_b, align=al_c)
    if isinstance(thua_gio, float) and thua_gio % 1 != 0:
        c_sec8.number_format = '0.0'

    # H. CHỮ KÝ VÀ NGÀY THÁNG
    r_date = r_sec8 + 2
    r_title = r_date + 1
    r_names = r_title + 5

    for r_cl in range(r_sec8 + 1, r_names + 4):
        for c_cl in range(1, 12):
            cell = ws.cell(row=r_cl, column=c_cl)
            if not isinstance(cell, MergedCell):
                cell.value = None

    font_sign = Font(name='Times New Roman', size=FONT_SIZE, bold=True)
    str_date = (
        f"Cần Thơ, ngày {D['d_end'].day:02d} tháng {D['d_end'].month:02d} năm"
        f" {D['d_end'].year}"
    )
    unmerge_and_merge_exact(
        ws, f'H{r_date}:K{r_date}', str_date, font=font_it, align=al_c
    )

    unmerge_and_merge_exact(
        ws, f'B{r_title}:D{r_title}', 'Phó Hiệu trưởng', font=font_sign, align=al_c
    )
    unmerge_and_merge_exact(
        ws, f'F{r_title}:G{r_title}', 'Tổ trưởng', font=font_sign, align=al_c
    )
    unmerge_and_merge_exact(
        ws, f'H{r_title}:K{r_title}', 'Người kê khai', font=font_sign, align=al_c
    )
    ws.row_dimensions[r_title].height = 26

    unmerge_and_merge_exact(
        ws, f'B{r_names}:D{r_names}', D['ten_admin'], font=font_sign, align=al_c
    )
    unmerge_and_merge_exact(
        ws, f'F{r_names}:G{r_names}', D['ten_totruong'], font=font_sign, align=al_c
    )
    unmerge_and_merge_exact(
        ws, f'H{r_names}:K{r_names}', tengv, font=font_sign, align=al_c
    )
    ws.row_dimensions[r_names].height = 26

    out_path = os.path.abspath(
        f"PKG_Da_Dien_{thang_str.replace(' ', '_')}_{magv}.xlsx"
    )
    wb.save(out_path)
    return out_path