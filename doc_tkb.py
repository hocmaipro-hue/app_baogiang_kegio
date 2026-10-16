# doc_tkb.py
import openpyxl
import psycopg2
from datetime import datetime, date
from config import DB_URI

def xoa_toan_bo_tkb():
    """Hàm xóa sạch toàn bộ bảng thoi_khoa_bieu trên database"""
    try:
        conn = psycopg2.connect(DB_URI)
        cur = conn.cursor()
        cur.execute("DELETE FROM thoi_khoa_bieu;")
        conn.commit()
        cur.close()
        conn.close()
        return True, "Đã xóa sạch toàn bộ thời khóa biểu trên server thành công!"
    except Exception as e:
        return False, f"Lỗi khi xóa dữ liệu: {str(e)}"

def kiem_tra_tuan_ton_tai(ap_dung_tu_tuan):
    """Kiểm tra xem tuần này đã có bao nhiêu tiết trên database"""
    try:
        conn = psycopg2.connect(DB_URI)
        cur = conn.cursor()
        cur.execute("SELECT COUNT(*) FROM thoi_khoa_bieu WHERE ap_dung_tu_tuan = %s;", (int(ap_dung_tu_tuan),))
        row = cur.fetchone()
        cur.close()
        conn.close()
        return row[0] if row else 0
    except Exception:
        return 0

def chuan_hoa_ngay(ngay_input):
    if not ngay_input:
        return None
    if isinstance(ngay_input, (datetime, date)):
        return ngay_input if isinstance(ngay_input, date) else ngay_input.date()
    s = str(ngay_input).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None

def nap_tkb_tu_excel(file_path, ap_dung_tu_tuan=1, ap_dung_tu_ngay=None):
    try:
        conn = psycopg2.connect(DB_URI)
        cur = conn.cursor()

        ngay_chuan = chuan_hoa_ngay(ap_dung_tu_ngay)
        ngay_hien_thi = ngay_chuan.strftime("%d/%m/%Y") if ngay_chuan else "Chưa xác định"

        # Tự động nhận diện cột ngày
        cur.execute("""
            SELECT column_name FROM information_schema.columns 
            WHERE table_name = 'thoi_khoa_bieu';
        """)
        existing_cols = [r[0] for r in cur.fetchall()]
        col_ngay = None
        if "ap_dung_từ_ngay" in existing_cols:
            col_ngay = '"ap_dung_từ_ngay"'
        elif "ap_dung_tu_ngay" in existing_cols:
            col_ngay = "ap_dung_tu_ngay"

        # Lấy danh sách GV
        cur.execute("SELECT magv, tengv FROM giao_vien;")
        ds_gv = cur.fetchall()
        map_magv, map_ten_tat = {}, {}
        for g in ds_gv:
            magv_db, tengv = str(g[0]).strip(), str(g[1]).strip()
            map_magv[magv_db.upper()] = magv_db
            map_ten_tat[tengv.upper()] = magv_db
            parts = tengv.split()
            if parts and parts[-1].upper() not in map_ten_tat:
                map_ten_tat[parts[-1].upper()] = magv_db

        # Đọc Excel
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb["DATASHEET"] if "DATASHEET" in wb.sheetnames else wb.active

        danh_sach_tkb = []
        header_rows = [3, 65, 126, 187]

        for i, header_row in enumerate(header_rows):
            if header_row > sheet.max_row:
                continue
            lop_columns = {}
            for col in range(4, sheet.max_column + 1):
                val = sheet.cell(row=header_row, column=col).value
                if val and not str(val).strip().isdigit():
                    lop_columns[col] = str(val).strip()

            if not lop_columns:
                continue

            start_row = header_row + 1
            end_row = header_rows[i + 1] - 1 if i + 1 < len(header_rows) else sheet.max_row

            current_thu, current_buoi = 2, "Sáng"

            for row in range(start_row, end_row + 1):
                val_thu = sheet.cell(row=row, column=1).value
                val_buoi = sheet.cell(row=row, column=2).value
                val_tiet = sheet.cell(row=row, column=3).value

                if val_thu:
                    digits = ''.join(filter(str.isdigit, str(val_thu)))
                    if digits: current_thu = int(digits)

                if val_buoi:
                    b_str = str(val_buoi).strip().upper()
                    if b_str in ['C', 'CHIỀU', 'CHIEU']: current_buoi = "Chiều"
                    elif b_str in ['S', 'SÁNG', 'SANG']: current_buoi = "Sáng"

                if not val_tiet or not str(val_tiet).strip().isdigit():
                    continue
                tiet_num = int(str(val_tiet).strip())

                for col, ten_lop in lop_columns.items():
                    cell_val = sheet.cell(row=row, column=col).value
                    if not cell_val: continue
                    s_cell = str(cell_val).strip()
                    if "-" in s_cell:
                        parts = s_cell.split("-")
                        mon = parts[0].strip()
                        raw_gv = parts[1].strip() if len(parts) > 1 else ""
                        magv_chuan = map_magv.get(raw_gv.upper(), map_ten_tat.get(raw_gv.upper(), raw_gv))

                        if col_ngay:
                            danh_sach_tkb.append((magv_chuan, current_thu, current_buoi, tiet_num, ten_lop, mon, int(ap_dung_tu_tuan), ngay_chuan))
                        else:
                            danh_sach_tkb.append((magv_chuan, current_thu, current_buoi, tiet_num, ten_lop, mon, int(ap_dung_tu_tuan)))

        if not danh_sach_tkb:
            cur.close()
            conn.close()
            return False, "Không tìm thấy dữ liệu tiết dạy hợp lệ trong file TKB!"

        # 👉 XÓA BẢN CŨ CỦA ĐÚNG TUẦN NÀY ĐỂ GHI ĐÈ SẠCH SẼ
        cur.execute("DELETE FROM thoi_khoa_bieu WHERE ap_dung_tu_tuan = %s;", (int(ap_dung_tu_tuan),))

        # 👉 INSERT BẢN MỚI
        if col_ngay:
            insert_query = f"""
                INSERT INTO thoi_khoa_bieu (magv, thu, buoi, tiet, lop, mon, ap_dung_tu_tuan, {col_ngay})
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            """
        else:
            insert_query = """
                INSERT INTO thoi_khoa_bieu (magv, thu, buoi, tiet, lop, mon, ap_dung_tu_tuan)
                VALUES (%s, %s, %s, %s, %s, %s, %s)
            """

        cur.executemany(insert_query, danh_sach_tkb)
        conn.commit()
        cur.close()
        conn.close()

        return True, f"Bóc tách thành công {len(danh_sach_tkb)} tiết học (Đã ghi đè vào Tuần {ap_dung_tu_tuan} - Từ ngày: {ngay_hien_thi})!"

    except Exception as e:
        return False, f"Lỗi đọc file TKB: {str(e)}"