# excel_importer.py
import openpyxl
import psycopg2
from config import DB_URI

def safe_float(val, default=0.0):
    """Hàm phụ trợ chuyển đổi an toàn sang kiểu số thực (float)"""
    try:
        return float(val) if val is not None and str(val).strip() != "" else default
    except ValueError:
        return default

def nhap_dsgv(file_path):
    """
    Đọc file Excel danh sách giáo viên bằng openpyxl và nạp toàn bộ thông tin cố định vào bảng giao_vien.
    Tự động cấp mật khẩu mặc định '123456' nếu chưa có.
    """
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active

        danh_sach_gv = []
        
        # Duyệt từ dòng 2 (bỏ qua dòng tiêu đề hàng 1)
        for row in range(2, sheet.max_row + 1):
            val_magv = sheet.cell(row=row, column=1).value
            val_tengv = sheet.cell(row=row, column=2).value
            
            # Bỏ qua dòng trống mã hoặc tên giáo viên
            if not val_magv or not val_tengv:
                continue

            magv_str = str(val_magv).strip()
            tengv_str = str(val_tengv).strip()
            to_str = str(sheet.cell(row=row, column=3).value or "").strip() or "Chung"
            role_str = str(sheet.cell(row=row, column=4).value or "").strip().upper() or "GV"
            chucvu_str = str(sheet.cell(row=row, column=5).value or "").strip() or "Giáo viên"
            kiemnhiem_str = str(sheet.cell(row=row, column=6).value or "").strip()
            
            # Các cột thông tin cố định mới
            tiet_giamtru = safe_float(sheet.cell(row=row, column=7).value, 0.0)
            con_nho_str = str(sheet.cell(row=row, column=8).value or "").strip()
            giam_tru_cn = safe_float(sheet.cell(row=row, column=9).value, 0.0)
            hd_lan_dau_str = str(sheet.cell(row=row, column=10).value or "").strip()
            giam_tru_hd = safe_float(sheet.cell(row=row, column=11).value, 0.0)
            kiem_nhiem_2_str = str(sheet.cell(row=row, column=12).value or "").strip()
            giam_tru_kn2 = safe_float(sheet.cell(row=row, column=13).value, 0.0)

            danh_sach_gv.append({
                "magv": magv_str,
                "tengv": tengv_str,
                "matkhau": "123456",  # Mật khẩu khởi tạo mặc định
                "to_bo_mon": to_str,
                "role": role_str,
                "chuc_vu": chucvu_str,
                "kiem_nhiem": kiemnhiem_str,
                "dinh_muc_tiet": 17,  # Định mức chuẩn mặc định
                "tiet_giam_tru": tiet_giamtru,
                "con_nho": con_nho_str,
                "giam_tru_con_nho": giam_tru_cn,
                "hopdong_lan_dau": hd_lan_dau_str,
                "giam_tru_hopdong_landau": giam_tru_hd,
                "kiem_nhiem_2": kiem_nhiem_2_str,
                "giam_tru_kn2": giam_tru_kn2
            })

        if not danh_sach_gv:
            return False, "File Excel không có dữ liệu giáo viên hợp lệ!"

        # Nạp dữ liệu vào PostgreSQL / Supabase bằng psycopg2
        connection = psycopg2.connect(DB_URI)
        cursor = connection.cursor()
        
        count = 0
        for gv in danh_sach_gv:
            sql = """
                INSERT INTO giao_vien (
                    magv, tengv, matkhau, to_bo_mon, role, chuc_vu, kiem_nhiem, 
                    dinh_muc_tiet, tiet_giam_tru, con_nho, giam_tru_con_nho, 
                    hopdong_lan_dau, giam_tru_hopdong_landau, kiem_nhiem_2, giam_tru_kn2
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (magv) DO UPDATE 
                SET tengv = EXCLUDED.tengv, 
                    to_bo_mon = EXCLUDED.to_bo_mon, 
                    role = EXCLUDED.role, 
                    chuc_vu = EXCLUDED.chuc_vu, 
                    kiem_nhiem = EXCLUDED.kiem_nhiem, 
                    tiet_giam_tru = EXCLUDED.tiet_giam_tru,
                    con_nho = EXCLUDED.con_nho,
                    giam_tru_con_nho = EXCLUDED.giam_tru_con_nho,
                    hopdong_lan_dau = EXCLUDED.hopdong_lan_dau,
                    giam_tru_hopdong_landau = EXCLUDED.giam_tru_hopdong_landau,
                    kiem_nhiem_2 = EXCLUDED.kiem_nhiem_2,
                    giam_tru_kn2 = EXCLUDED.giam_tru_kn2;
            """
            cursor.execute(sql, (
                gv["magv"], gv["tengv"], gv["matkhau"], gv["to_bo_mon"], 
                gv["role"], gv["chuc_vu"], gv["kiem_nhiem"], gv["dinh_muc_tiet"], 
                gv["tiet_giam_tru"], gv["con_nho"], gv["giam_tru_con_nho"], 
                gv["hopdong_lan_dau"], gv["giam_tru_hopdong_landau"], 
                gv["kiem_nhiem_2"], gv["giam_tru_kn2"]
            ))
            count += 1
            
        connection.commit()
        cursor.close()
        connection.close()

        return True, f"Đã nạp thành công {count} giáo viên cùng toàn bộ thông tin cố định vào cơ sở dữ liệu!"

    except Exception as e:
        return False, f"Lỗi đọc file Excel: {str(e)}"
